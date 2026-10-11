"""
Format-Preserving PowerPoint (.pptx) Redaction Engine.
Redacts slide titles, text boxes, tables, speaker notes, alternate text, core metadata,
and embedded images while preserving slide layouts, themes, and unaffected formatting.
"""
from __future__ import annotations

import io
import zipfile
from typing import Dict, List, Set, Tuple
import pptx

from app.models import BoundingBox, PIIFinding, ReviewStatus
from app.redactors.docx_redactor import replace_spans_across_runs
from app.redactors.image_redactor import redact_image_bytes_with_boxes


def redact_pptx_document(content: bytes, findings: List[PIIFinding]) -> bytes:
    """
    Apply format-preserving redaction to a PowerPoint (.pptx) presentation:
    1. Group active findings by block_id.
    2. Apply run-spanning replacement on slide text frames, table cells, and speaker notes.
    3. Sanitize shape alternate text (descr) and presentation core_properties.
    4. Pixel-redact any embedded media images where OCR detected PII.
    """
    prs = pptx.Presentation(io.BytesIO(content))

    by_block: Dict[str, List[Tuple[int, int, str]]] = {}
    by_image: Dict[str, List[Tuple[BoundingBox, str]]] = {}

    for f in findings:
        if f.review_status == ReviewStatus.REJECTED_FALSE_POSITIVE:
            continue
        if f.source_type == "ocr_embedded_image" and f.image_ref and f.bbox:
            by_image.setdefault(f.image_ref, []).append((f.bbox, f.replacement_label))
        elif f.source_type != "metadata":
            by_block.setdefault(f.block_id, []).append(
                (f.char_start, f.char_end, f.replacement_label)
            )

    for s_idx, slide in enumerate(prs.slides):
        slide_no = s_idx + 1
        for sh_idx, shape in enumerate(slide.shapes):
            # Alternate text sanitization
            alt_bid = f"pptx_s{slide_no}_sh{sh_idx}_alt"
            if alt_bid in by_block:
                try:
                    cNvPr = shape._element.xpath(".//p:cNvPr")
                    if cNvPr:
                        cNvPr[0].set("descr", "[ALT TEXT REDACTED]")
                except Exception:
                    pass

            if shape.has_text_frame:
                for p_idx, para in enumerate(shape.text_frame.paragraphs):
                    bid = f"pptx_s{slide_no}_sh{sh_idx}_p{p_idx}"
                    if bid in by_block:
                        replace_spans_across_runs(para.runs, by_block[bid], para)

            if shape.has_table:
                tbl = shape.table
                for r_idx, row in enumerate(tbl.rows):
                    for c_idx, cell in enumerate(row.cells):
                        for cp_idx, cp_para in enumerate(cell.text_frame.paragraphs):
                            bid = f"pptx_s{slide_no}_sh{sh_idx}_r{r_idx}_c{c_idx}_p{cp_idx}"
                            if bid in by_block:
                                replace_spans_across_runs(cp_para.runs, by_block[bid], cp_para)

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            for np_idx, npara in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                bid = f"pptx_s{slide_no}_note_p{np_idx}"
                if bid in by_block:
                    replace_spans_across_runs(npara.runs, by_block[bid], npara)

    # Sanitize Core Properties Metadata
    cp = prs.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""

    intermediate_buf = io.BytesIO()
    prs.save(intermediate_buf)
    pptx_bytes = intermediate_buf.getvalue()

    if not by_image:
        return pptx_bytes

    out_zip_buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(pptx_bytes), "r") as zin:
        with zipfile.ZipFile(out_zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                raw_data = zin.read(item.filename)
                if item.filename in by_image:
                    ext = "JPG" if item.filename.lower().endswith((".jpg", ".jpeg")) else "PNG"
                    raw_data = redact_image_bytes_with_boxes(
                        image_bytes=raw_data,
                        boxes_and_labels=by_image[item.filename],
                        output_format=ext,
                    )
                zout.writestr(item, raw_data)

    return out_zip_buf.getvalue()
