"""Format-aware PowerPoint (.pptx) Extractor handling slides, shapes, tables, speaker notes, alt text, metadata, and embedded images."""
from __future__ import annotations

import io
import zipfile
from typing import List, Set, Tuple
import pptx
from PIL import Image

from app.config import get_policy
from app.extractors.ocr_engine import run_ocr_on_image_bytes
from app.models import ExtractionProfile, TextBlockSegment


def extract_pptx_document(content: bytes) -> Tuple[List[TextBlockSegment], ExtractionProfile]:
    """
    Extract structured text segments from a PowerPoint (.pptx) presentation:
    - Core properties metadata
    - Slide titles, text frames, and shapes (concatenating runs per paragraph)
    - Slide tables and cells
    - Speaker notes
    - Shape alternate text (descr attribute)
    - Embedded raster images in ppt/media/* via OCR
    """
    policy = get_policy()
    prs = pptx.Presentation(io.BytesIO(content))
    segments: List[TextBlockSegment] = []

    metadata_count = 0
    paragraphs_count = 0
    tables_count = 0
    embedded_images_count = 0
    ocr_images_with_text = 0

    # 1. Core properties metadata
    cp = prs.core_properties
    meta_map = {
        "author": cp.author,
        "last_modified_by": cp.last_modified_by,
        "title": cp.title,
        "subject": cp.subject,
        "comments": cp.comments,
    }
    for k, val in meta_map.items():
        if val and str(val).strip():
            metadata_count += 1
            segments.append(
                TextBlockSegment(
                    block_id=f"pptx_meta_{k}",
                    source_type="metadata",
                    page_or_slide=1,
                    location_ref=f"PPTX Metadata · {k}",
                    text=str(val).strip(),
                    metadata_key=k,
                )
            )

    # 2. Slide-by-slide extraction
    for s_idx, slide in enumerate(prs.slides):
        slide_no = s_idx + 1
        title_shape = slide.shapes.title

        for sh_idx, shape in enumerate(slide.shapes):
            # Alternate text check
            try:
                cNvPr = shape._element.xpath(".//p:cNvPr")
                if cNvPr:
                    descr = (cNvPr[0].get("descr") or "").strip()
                    if descr:
                        segments.append(
                            TextBlockSegment(
                                block_id=f"pptx_s{slide_no}_sh{sh_idx}_alt",
                                source_type="alt_text",
                                page_or_slide=slide_no,
                                location_ref=f"Slide {slide_no} · Shape '{shape.name}' Alt-Text",
                                text=descr,
                            )
                        )
            except Exception:
                pass

            if shape.has_text_frame:
                is_title = title_shape is not None and shape == title_shape
                for p_idx, para in enumerate(shape.text_frame.paragraphs):
                    txt = para.text.strip()
                    if not txt:
                        continue
                    paragraphs_count += 1
                    stype = "slide_title" if (is_title or (sh_idx == 0 and p_idx == 0)) else "slide_text"
                    segments.append(
                        TextBlockSegment(
                            block_id=f"pptx_s{slide_no}_sh{sh_idx}_p{p_idx}",
                            source_type=stype,  # type: ignore[arg-type]
                            page_or_slide=slide_no,
                            location_ref=f"Slide {slide_no} · {shape.name} · P{p_idx + 1}",
                            text=txt,
                        )
                    )

            if shape.has_table:
                tables_count += 1
                tbl = shape.table
                for r_idx, row in enumerate(tbl.rows):
                    for c_idx, cell in enumerate(row.cells):
                        for cp_idx, cp_para in enumerate(cell.text_frame.paragraphs):
                            ctxt = cp_para.text.strip()
                            if not ctxt:
                                continue
                            segments.append(
                                TextBlockSegment(
                                    block_id=f"pptx_s{slide_no}_sh{sh_idx}_r{r_idx}_c{c_idx}_p{cp_idx}",
                                    source_type="table_cell",
                                    page_or_slide=slide_no,
                                    location_ref=f"Slide {slide_no} · Table '{shape.name}' · R{r_idx + 1}C{c_idx + 1}",
                                    text=ctxt,
                                    table_coords={
                                        "slide": s_idx,
                                        "shape": sh_idx,
                                        "row": r_idx,
                                        "col": c_idx,
                                        "para": cp_idx,
                                    },
                                )
                            )

        # Speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            for np_idx, npara in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                ntxt = npara.text.strip()
                # Ignore default slide-number placeholders in notes
                if ntxt and not ntxt.isdigit():
                    segments.append(
                        TextBlockSegment(
                            block_id=f"pptx_s{slide_no}_note_p{np_idx}",
                            source_type="speaker_note",
                            page_or_slide=slide_no,
                            location_ref=f"Slide {slide_no} · Speaker Notes · P{np_idx + 1}",
                            text=ntxt,
                        )
                    )

    # 3. Embedded images in ppt/media/* via OCR
    with zipfile.ZipFile(io.BytesIO(content), "r") as zf:
        media_files = sorted(
            n
            for n in zf.namelist()
            if n.startswith("ppt/media/")
            and n.lower().endswith((".png", ".jpg", ".jpeg"))
        )
        embedded_images_count = len(media_files)
        if policy.enable_ocr_on_embedded_images:
            for img_idx, m_name in enumerate(media_files):
                img_bytes = zf.read(m_name)
                try:
                    im = Image.open(io.BytesIO(img_bytes))
                    # Skip small logos/icons that do not contain document content
                    if im.width < 120 or im.height < 55:
                        continue
                except Exception:
                    continue

                ocr_segs = run_ocr_on_image_bytes(
                    image_bytes=img_bytes,
                    page_or_slide=1,
                    location_prefix=f"PPTX Embedded Image ({m_name})",
                    source_type="ocr_embedded_image",
                    image_ref=m_name,
                    block_id_prefix=f"pptx_img{img_idx + 1}",
                )
                if ocr_segs:
                    ocr_images_with_text += 1
                    segments.extend(ocr_segs)

    notes = [
        f"Processed {len(prs.slides)} slides, {paragraphs_count} text blocks, and {tables_count} slide tables.",
    ]
    if embedded_images_count > 0:
        notes.append(
            f"Inspected {embedded_images_count} media asset(s); {ocr_images_with_text} contained OCR text."
        )

    profile = ExtractionProfile(
        detected_format="PPTX",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        is_scanned_document=False,
        total_pages_or_slides=len(prs.slides),
        native_text_pages_list=list(range(1, len(prs.slides) + 1)),
        tables_count=tables_count,
        paragraphs_count=paragraphs_count,
        embedded_images_count=embedded_images_count,
        ocr_embedded_images_with_text=ocr_images_with_text,
        metadata_fields_extracted=metadata_count,
        extraction_notes=notes,
    )
    return segments, profile
