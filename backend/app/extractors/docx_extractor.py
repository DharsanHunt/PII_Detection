"""Format-aware Word (.docx) Extractor handling paragraphs, runs, tables, headers/footers, metadata, and embedded image OCR."""
from __future__ import annotations

import io
import zipfile
from typing import List, Set, Tuple
import docx
from PIL import Image

from app.config import get_policy
from app.extractors.ocr_engine import run_ocr_on_image_bytes
from app.models import ExtractionProfile, TextBlockSegment


def extract_docx_document(content: bytes) -> Tuple[List[TextBlockSegment], ExtractionProfile]:
    """
    Extract structured text segments from a Word (.docx) package:
    - Core properties metadata
    - Section headers and footers
    - Body paragraphs (preserving concatenated run text for cross-run entity detection)
    - Table cells (deduplicating merged cells)
    - Embedded raster screenshots in word/media/* via OCR
    """
    policy = get_policy()
    doc = docx.Document(io.BytesIO(content))
    segments: List[TextBlockSegment] = []

    metadata_count = 0
    paragraphs_count = 0
    tables_count = len(doc.tables)
    embedded_images_count = 0
    ocr_images_with_text = 0

    # 1. Core properties metadata
    cp = doc.core_properties
    meta_map = {
        "author": cp.author,
        "last_modified_by": cp.last_modified_by,
        "title": cp.title,
        "subject": cp.subject,
        "comments": cp.comments,
    }
    for k, val in meta_map.items():
        if val and str(val).strip():
            metadata_count += 1
            segments.append(
                TextBlockSegment(
                    block_id=f"docx_meta_{k}",
                    source_type="metadata",
                    page_or_slide=1,
                    location_ref=f"DOCX Metadata · {k}",
                    text=str(val).strip(),
                    metadata_key=k,
                )
            )

    # 2. Section headers and footers
    seen_header_footer_texts: Set[str] = set()
    for s_idx, section in enumerate(doc.sections):
        if section.header:
            for p_idx, p in enumerate(section.header.paragraphs):
                txt = p.text.strip()
                if txt and txt not in seen_header_footer_texts:
                    seen_header_footer_texts.add(txt)
                    segments.append(
                        TextBlockSegment(
                            block_id=f"docx_sec{s_idx}_hdr_p{p_idx}",
                            source_type="header",
                            page_or_slide=1,
                            location_ref=f"Section {s_idx + 1} · Header · P{p_idx + 1}",
                            text=txt,
                        )
                    )
        if section.footer:
            for p_idx, p in enumerate(section.footer.paragraphs):
                txt = p.text.strip()
                if txt and txt not in seen_header_footer_texts:
                    seen_header_footer_texts.add(txt)
                    segments.append(
                        TextBlockSegment(
                            block_id=f"docx_sec{s_idx}_ftr_p{p_idx}",
                            source_type="footer",
                            page_or_slide=1,
                            location_ref=f"Section {s_idx + 1} · Footer · P{p_idx + 1}",
                            text=txt,
                        )
                    )

    # 3. Body paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        paragraphs_count += 1
        style_name = p.style.name if p.style and p.style.name else "Normal"
        est_page = max(1, (paragraphs_count // 28) + 1)
        segments.append(
            TextBlockSegment(
                block_id=f"docx_p_{p_idx}",
                source_type="paragraph",
                page_or_slide=est_page,
                location_ref=f"Paragraph {p_idx + 1} ({style_name})",
                text=txt,
            )
        )

    # 4. Tables and cells (deduplicating merged cells sharing underlying _tc element)
    for t_idx, table in enumerate(doc.tables):
        seen_tcs: List[object] = []
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                tc_obj = cell._tc
                if any(tc_obj is existing for existing in seen_tcs):
                    continue
                seen_tcs.append(tc_obj)
                for cp_idx, cp_para in enumerate(cell.paragraphs):
                    ctxt = cp_para.text.strip()
                    if not ctxt:
                        continue
                    segments.append(
                        TextBlockSegment(
                            block_id=f"docx_t{t_idx}_r{r_idx}_c{c_idx}_p{cp_idx}",
                            source_type="table_cell",
                            page_or_slide=1,
                            location_ref=f"Table {t_idx + 1} · Row {r_idx + 1} · Col {c_idx + 1}",
                            text=ctxt,
                            table_coords={
                                "table": t_idx,
                                "row": r_idx,
                                "col": c_idx,
                                "para": cp_idx,
                            },
                        )
                    )

    # 5. Embedded images in word/media/* via OCR
    with zipfile.ZipFile(io.BytesIO(content), "r") as zf:
        media_files = sorted(
            n
            for n in zf.namelist()
            if n.startswith("word/media/")
            and n.lower().endswith((".png", ".jpg", ".jpeg"))
        )
        embedded_images_count = len(media_files)
        if policy.enable_ocr_on_embedded_images:
            for img_idx, m_name in enumerate(media_files):
                img_bytes = zf.read(m_name)
                try:
                    im = Image.open(io.BytesIO(img_bytes))
                    if im.width < 90 or im.height < 40:
                        continue
                except Exception:
                    continue

                short_ref = m_name.replace("word/", "")
                ocr_segs = run_ocr_on_image_bytes(
                    image_bytes=img_bytes,
                    page_or_slide=1,
                    location_prefix=f"DOCX Embedded Screenshot #{img_idx + 1}",
                    source_type="ocr_embedded_image",
                    image_ref=m_name,
                    block_id_prefix=f"docx_img{img_idx + 1}",
                )
                if ocr_segs:
                    ocr_images_with_text += 1
                    segments.extend(ocr_segs)

    notes = [
        f"Extracted {paragraphs_count} non-empty paragraphs and {tables_count} structured tables.",
    ]
    if embedded_images_count > 0:
        notes.append(
            f"Inspected {embedded_images_count} embedded media images; {ocr_images_with_text} contained OCR-recoverable text."
        )

    profile = ExtractionProfile(
        detected_format="DOCX",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        is_scanned_document=False,
        total_pages_or_slides=max(1, (paragraphs_count // 28) + 1),
        native_text_pages_list=list(range(1, max(1, (paragraphs_count // 28) + 1) + 1)),
        tables_count=tables_count,
        paragraphs_count=paragraphs_count,
        embedded_images_count=embedded_images_count,
        ocr_embedded_images_with_text=ocr_images_with_text,
        metadata_fields_extracted=metadata_count,
        extraction_notes=notes,
    )
    return segments, profile
